import sys
import os
import base64
import secrets
from pathlib import Path

from flask import Flask, render_template, request, jsonify, Blueprint, flash
from flask_login import current_user, login_required
from werkzeug.utils import secure_filename

WEB_DIR = Path(__file__).resolve().parent
sys.path.insert(0, str(WEB_DIR.parent))

from aes_core import (
    SBOX, INV_SBOX, RCON,
    parse_hex_block, parse_hex_key, key_expansion,
    encrypt_block, decrypt_block,
    encrypt_block_with_rounds, decrypt_block_with_rounds,
    one_round_steps, words_to_hex_words, state_to_hex_grid, state_to_bytes, bit_mismatch,
)
from extensions import db, bcrypt, login_manager
from models import CoachAttempt, EncryptionAttempt

try:
    from docx import Document as DocxDocument
    DOCX_AVAILABLE = True
except ImportError:
    DOCX_AVAILABLE = False

KEY_BLOCK_LEN = 32
DEFAULT_KEY   = "2b7e151628aed2a6abf7158809cf4f3c"
DEFAULT_BLOCK = "3243f6a8885a308d313198a2e0370734"
NIST_KEY      = DEFAULT_KEY
NIST_PT       = DEFAULT_BLOCK
NIST_CT       = "3925841d02dc09fbdc118597196a0b32"

def normalize_hex(s):
    if not s: return ""
    return s.replace(" ", "").replace("\n", "").replace("\t", "").strip().lower()

def is_valid_hex(s, length=32):
    n = normalize_hex(s)
    return len(n) == length and all(c in "0123456789abcdef" for c in n)

def _validate_hex_chars(n, label):
    if any(c not in "0123456789abcdef" for c in n):
        raise ValueError(f"{label} can only contain hex characters: 0-9 and a-f.")

def pkcs7_pad(data):
    pad_len = 16 - (len(data) % 16)
    return data + bytes([pad_len]) * pad_len

def pkcs7_unpad(data):
    if not data: raise ValueError("Data is empty.")
    pad_len = data[-1]
    if pad_len < 1 or pad_len > 16: raise ValueError("Invalid PKCS#7 padding.")
    if data[-pad_len:] != bytes([pad_len]) * pad_len: raise ValueError("PKCS#7 padding bytes are inconsistent.")
    return data[:-pad_len]

def coerce_key_input(value, label):
    n = normalize_hex(value)
    if not n: raise ValueError(f"{label} is required (up to 32 hex characters).")
    _validate_hex_chars(n, label)
    if len(n) > KEY_BLOCK_LEN: raise ValueError(f"{label} is too long ({len(n)} hex chars). Maximum is 32.")
    return n.ljust(KEY_BLOCK_LEN, "0")

def coerce_hex_input(value, label):
    n = normalize_hex(value)
    if not n: raise ValueError(f"{label} is required.")
    _validate_hex_chars(n, label)
    if len(n) % 2 != 0: n += "0"
    raw = bytes.fromhex(n)
    if len(raw) > 16: raw = raw[:16]
    elif len(raw) < 16: raw = pkcs7_pad(raw)
    return raw.hex()

def hex_to_padded_blocks(value, label):
    n = normalize_hex(value)
    if not n: raise ValueError(f"{label} is required.")
    _validate_hex_chars(n, label)
    if len(n) % 2 != 0: n += "0"
    raw = bytes.fromhex(n)
    if len(raw) % 16 != 0: raw = pkcs7_pad(raw)
    return [raw[i:i+16] for i in range(0, len(raw), 16)]

def compare_hex_answer(answer, expected_hex, label, required=False):
    raw = answer or ""
    normalized = normalize_hex(raw)
    if not normalized:
        if required: return {"label": label, "answer": "", "expected": expected_hex, "status": "wrong", "message": "Missing answer"}
        return {"label": label, "answer": "", "expected": expected_hex, "status": "skipped", "message": "Not submitted"}
    if len(normalized) != KEY_BLOCK_LEN or any(c not in "0123456789abcdef" for c in normalized):
        return {"label": label, "answer": normalized, "expected": expected_hex, "status": "wrong", "message": "Use exactly 32 hex characters"}
    if normalized == expected_hex:
        return {"label": label, "answer": normalized, "expected": expected_hex, "status": "correct", "message": "Correct"}
    return {"label": label, "answer": normalized, "expected": expected_hex, "status": "wrong", "message": "Different from expected state"}

def byte_bit_mismatch(a, b):
    return sum((left ^ right).bit_count() for left, right in zip(a, b))

def build_coach_practice(mode, key, block, final_answer, round_answers):
    key   = coerce_key_input(key, "Key")
    block = coerce_hex_input(block, "Block")
    if mode not in ("encrypt", "decrypt"): mode = "encrypt"
    key_b = parse_hex_key(key); block_b = parse_hex_block(block); rks = key_expansion(key_b)
    if mode == "encrypt":
        expected_output = encrypt_block(block_b, rks).hex()
        states = encrypt_block_with_rounds(block_b, rks)
        input_label, output_label, round_hint = "Plaintext", "Ciphertext", "Encryption round boundary"
    else:
        expected_output = decrypt_block(block_b, rks).hex()
        states = decrypt_block_with_rounds(block_b, rks)
        input_label, output_label, round_hint = "Ciphertext", "Plaintext", "Decryption inverse-round boundary"
    final_check = compare_hex_answer(final_answer, expected_output, output_label, required=True)
    step_checks = []
    for rnd, state in enumerate(states):
        check = compare_hex_answer(round_answers[rnd], state_to_bytes(state).hex(), f"After round {rnd}")
        check["grid"] = state_to_hex_grid(state)
        check["hint"] = round_hint
        step_checks.append(check)
    answered = [s for s in step_checks if s["status"] != "skipped"]
    correct  = [s for s in answered  if s["status"] == "correct"]
    return {"mode": mode, "input_label": input_label, "output_label": output_label,
            "key_hex": key_b.hex(), "block_hex": block_b.hex(), "expected_output": expected_output,
            "final_check": final_check, "step_checks": step_checks,
            "answered_count": len(answered), "correct_count": len(correct)}

def rot_word(word): return ((word << 8) | (word >> 24)) & 0xFFFFFFFF

def sub_word(word):
    return (SBOX[(word >> 24) & 0xFF] << 24 | SBOX[(word >> 16) & 0xFF] << 16 |
            SBOX[(word >> 8)  & 0xFF] << 8  | SBOX[word & 0xFF])

def key_expansion_word_details(key_b):
    words = [int.from_bytes(key_b[i*4:(i+1)*4], "big") for i in range(4)]
    details = [{"index": i, "round": 0, "word": f"{words[i]:08x}",
                "formula": f"w[{i}] = original key word {i}", "special": False} for i in range(4)]
    for i in range(4, 44):
        previous = words[i-1]; temp = previous
        step = {"index": i, "round": i//4, "previous": f"{previous:08x}",
                "from_four_back": f"{words[i-4]:08x}", "special": i % 4 == 0}
        if i % 4 == 0:
            rotated = rot_word(previous); substituted = sub_word(rotated); rcon_word = RCON[i//4] << 24
            temp = substituted ^ rcon_word
            step.update({"rotated": f"{rotated:08x}", "substituted": f"{substituted:08x}",
                          "rcon": f"{rcon_word:08x}", "temp": f"{temp:08x}",
                          "formula": f"w[{i}] = w[{i-4}] XOR SubWord(RotWord(w[{i-1}])) XOR Rcon[{i//4}]"})
        else:
            step.update({"temp": f"{temp:08x}", "formula": f"w[{i}] = w[{i-4}] XOR w[{i-1}]"})
        new_word = words[i-4] ^ temp; words.append(new_word); step["word"] = f"{new_word:08x}"
        details.append(step)
    return details

def xor_blocks(left, right): return bytes(a ^ b for a, b in zip(left, right))

def aes_cbc_encrypt(plaintext_bytes, key_b):
    rks = key_expansion(key_b); iv = secrets.token_bytes(16); prev = iv; ct = bytearray()
    for i in range(0, len(plaintext_bytes), 16):
        blk = plaintext_bytes[i:i+16]; enc = encrypt_block(xor_blocks(blk, prev), rks)
        ct.extend(enc); prev = enc
    return iv + bytes(ct)

def aes_cbc_decrypt(iv_plus_ct, key_b):
    iv = iv_plus_ct[:16]; ct = iv_plus_ct[16:]; rks = key_expansion(key_b); prev = iv; pt = bytearray()
    for i in range(0, len(ct), 16):
        blk = ct[i:i+16]; dec = decrypt_block(blk, rks); pt.extend(xor_blocks(dec, prev)); prev = blk
    return bytes(pt)

def extract_docx_text(data):
    import io; doc = DocxDocument(io.BytesIO(data))
    return "\n".join(p.text for p in doc.paragraphs if p.text.strip())

def make_docx_with_text(text):
    import io; doc = DocxDocument()
    for i in range(0, max(len(text), 1), 80): doc.add_paragraph(text[i:i+80])
    buf = io.BytesIO(); doc.save(buf); return buf.getvalue()

def format_file_size(size):
    if size < 1024: return f"{size} bytes"
    if size < 1024*1024: return f"{size/1024:.1f} KB"
    return f"{size/(1024*1024):.1f} MB"

def build_doc_crypto_result(mode, key, storage):
    key = coerce_key_input(key, "Key"); key_b = parse_hex_key(key)
    if not storage or not storage.filename: raise ValueError("Please choose a file to upload.")
    filename = secure_filename(storage.filename) or "document"; data = storage.read()
    if not data: raise ValueError("The uploaded file is empty.")
    if mode == "encrypt":
        lower = filename.lower()
        if lower.endswith(".docx"):
            if not DOCX_AVAILABLE: raise ValueError("python-docx not installed.")
            plaintext_str = extract_docx_text(data)
        elif lower.endswith(".txt"):
            plaintext_str = data.decode("utf-8", errors="replace")
        else:
            raise ValueError("Upload a .docx or .txt file to encrypt.")
        if not plaintext_str.strip(): raise ValueError("The document is empty — nothing to encrypt.")
        padded = pkcs7_pad(plaintext_str.encode("utf-8")); iv_ct = aes_cbc_encrypt(padded, key_b)
        ct_hex = iv_ct.hex(); out_bytes = make_docx_with_text(ct_hex)
        out_name = filename.rsplit(".", 1)[0] + "_encrypted.docx"
        label = "Encrypted document (ciphertext as hex inside .docx)"
        preview_p = plaintext_str[:300] + ("…" if len(plaintext_str) > 300 else "")
        preview_c = ct_hex[:160] + ("…" if len(ct_hex) > 160 else "")
    else:
        if not filename.lower().endswith(".docx"): raise ValueError("Upload the encrypted .docx produced by this tool.")
        if not DOCX_AVAILABLE: raise ValueError("python-docx not installed.")
        ct_hex_raw = extract_docx_text(data); ct_hex = normalize_hex(ct_hex_raw)
        if not ct_hex or any(c not in "0123456789abcdef" for c in ct_hex): raise ValueError("File does not contain valid AES ciphertext.")
        if len(ct_hex) < 64 or len(ct_hex) % 2 != 0: raise ValueError("Ciphertext is too short or malformed.")
        iv_ct = bytes.fromhex(ct_hex); padded_pt = aes_cbc_decrypt(iv_ct, key_b)
        pt_bytes = pkcs7_unpad(padded_pt); plaintext_str = pt_bytes.decode("utf-8", errors="replace")
        out_bytes = make_docx_with_text(plaintext_str)
        out_name = filename.replace("_encrypted", "").rsplit(".", 1)[0] + "_decrypted.docx"
        label = "Decrypted document (plaintext)"
        preview_p = plaintext_str[:300] + ("…" if len(plaintext_str) > 300 else "")
        preview_c = ct_hex[:160] + ("…" if len(ct_hex) > 160 else "")
    return {
        "mode": mode, "label": label,
        "input_name": filename, "input_size": format_file_size(len(data)),
        "output_name": out_name, "output_size": format_file_size(len(out_bytes)),
        "key_hex": key_b.hex(),
        "download_b64":  base64.b64encode(out_bytes).decode("ascii"),
        "download_mime": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        "preview_plain":  preview_p,
        "preview_cipher": preview_c,
    }

# ── APP FACTORY ──────────────────────────────────────────────────────────────
def create_app():
    app = Flask(__name__, template_folder=str(WEB_DIR / "templates"),
                static_folder=str(WEB_DIR / "static"))
    app.config["SECRET_KEY"] = os.environ.get("SECRET_KEY", secrets.token_hex(32))
    app.config["MAX_CONTENT_LENGTH"] = 8 * 1024 * 1024
    db_url = os.environ.get("DATABASE_URL", "sqlite:///aes_app.db")
    if db_url.startswith("postgres://"): db_url = db_url.replace("postgres://", "postgresql://", 1)
    app.config["SQLALCHEMY_DATABASE_URI"] = db_url
    app.config["SQLALCHEMY_TRACK_MODIFICATIONS"] = False
    db.init_app(app); bcrypt.init_app(app); login_manager.init_app(app)
    from auth import auth as auth_bp
    app.register_blueprint(auth_bp)
    app.register_blueprint(main_bp)
    with app.app_context(): db.create_all()
    return app

# ── BLUEPRINT — must be defined before routes ────────────────────────────────
main_bp = Blueprint("main", __name__)

# ── ROUTES ───────────────────────────────────────────────────────────────────
@main_bp.route("/")
def index():
    return render_template("index.html")

@main_bp.route("/single-round", methods=["GET", "POST"])
def single_round():
    error = None; steps_data = None
    key = request.form.get("key", DEFAULT_KEY).strip()
    block = request.form.get("block", DEFAULT_BLOCK).strip()
    round_idx = request.form.get("round", "1")
    saved = False
    if request.method == "POST":
        try:
            round_idx = int(round_idx)
            if not (0 <= round_idx <= 10): raise ValueError("Round must be 0–10.")
            key = coerce_key_input(key, "Key"); block = coerce_hex_input(block, "Block")
            key_b = parse_hex_key(key); block_b = parse_hex_block(block); rks = key_expansion(key_b)
            steps = one_round_steps(block_b, rks, round_idx)
            steps_data = [(name, state_to_hex_grid(state)) for name, state in steps]
            if current_user.is_authenticated:
                db.session.add(EncryptionAttempt(user_id=current_user.id, tool_name="single_round",
                    mode="encrypt", key_hex=key, input_hex=block))
                db.session.commit(); saved = True
        except Exception as e:
            error = str(e)
    try: round_idx = int(round_idx)
    except: round_idx = 1
    return render_template("single_round.html", key=key, block=block, round_idx=round_idx,
        steps_data=steps_data, error=error, saved=saved,
        key_hint="Up to 32 hex chars",
        block_hint="Any length — uses first 16 bytes, PKCS#7-padded if short")

@main_bp.route("/full-aes", methods=["GET", "POST"])
def full_aes():
    error = None; result = None; saved = False
    key = request.form.get("key", DEFAULT_KEY).strip()
    block = request.form.get("block", DEFAULT_BLOCK).strip()
    mode = request.form.get("mode", "encrypt")
    if request.method == "POST":
        try:
            key = coerce_key_input(key, "Key"); key_b = parse_hex_key(key); rks = key_expansion(key_b)
            blocks = hex_to_padded_blocks(block, "Input")
            raw_n = normalize_hex(block)
            if len(raw_n) % 2 != 0: raw_n += "0"
            original_byte_len = len(bytes.fromhex(raw_n))
            block_results = []
            if mode == "encrypt":
                for blk in blocks:
                    ct_b = encrypt_block(blk, rks)
                    block_results.append({"input_hex": blk.hex(), "output_hex": ct_b.hex(),
                                          "rounds": [state_to_hex_grid(s) for s in encrypt_block_with_rounds(blk, rks)]})
                result = {"mode": "encrypt", "input_label": "Plaintext", "output_label": "Ciphertext",
                          "input_hex": "".join(b["input_hex"] for b in block_results),
                          "output_hex": "".join(b["output_hex"] for b in block_results),
                          "block_results": block_results, "num_blocks": len(blocks),
                          "original_byte_len": original_byte_len, "rounds": block_results[0]["rounds"]}
            else:
                for blk in blocks:
                    pt_b = decrypt_block(blk, rks)
                    block_results.append({"input_hex": blk.hex(), "output_hex": pt_b.hex(),
                                          "rounds": [state_to_hex_grid(s) for s in decrypt_block_with_rounds(blk, rks)]})
                combined = bytes.fromhex("".join(b["output_hex"] for b in block_results))
                try: combined = pkcs7_unpad(combined)
                except ValueError: pass
                result = {"mode": "decrypt", "input_label": "Ciphertext", "output_label": "Plaintext",
                          "input_hex": "".join(b["input_hex"] for b in block_results),
                          "output_hex": combined.hex(),
                          "block_results": block_results, "num_blocks": len(blocks),
                          "original_byte_len": original_byte_len, "rounds": block_results[0]["rounds"]}
            if current_user.is_authenticated:
                db.session.add(EncryptionAttempt(user_id=current_user.id, tool_name="full_aes",
                    mode=mode, key_hex=key, input_hex=result["input_hex"], output_hex=result["output_hex"]))
                db.session.commit(); saved = True
        except Exception as e:
            error = str(e)
    return render_template("full_aes.html", key=key, block=block, mode=mode,
        result=result, error=error, saved=saved,
        key_hint="Up to 32 hex chars",
        block_hint="Any length — split into 16-byte blocks, last block PKCS#7-padded if needed")

@main_bp.route("/word-doc", methods=["GET", "POST"])
def word_doc():
    error = None; result = None; saved = False
    key = request.form.get("key", DEFAULT_KEY).strip()
    mode = request.form.get("mode", "encrypt")
    if request.method == "POST":
        try:
            if mode not in ("encrypt", "decrypt"): mode = "encrypt"
            result = build_doc_crypto_result(mode, key, request.files.get("document"))
            key = result["key_hex"]
            if current_user.is_authenticated:
                db.session.add(EncryptionAttempt(user_id=current_user.id, tool_name="file_crypt",
                    mode=mode, key_hex=result["key_hex"]))
                db.session.commit(); saved = True
        except Exception as e:
            error = str(e)
    return render_template("word_doc.html", key=key, mode=mode, result=result, error=error, saved=saved)

@main_bp.route("/avalanche", methods=["GET", "POST"])
def avalanche():
    error = None; result = None
    key1 = request.form.get("key1", DEFAULT_KEY).strip()
    block1 = request.form.get("block1", DEFAULT_BLOCK).strip()
    key2 = request.form.get("key2", DEFAULT_KEY).strip()
    block2 = request.form.get("block2", "3243f6a8885a308d313198a2e0370733").strip()
    mode = request.form.get("mode", "same_key")
    if request.method == "POST":
        try:
            if mode == "same_key":
                key1 = coerce_key_input(key1, "Key 1"); block1 = coerce_hex_input(block1, "Block 1")
                block2 = coerce_hex_input(block2, "Block 2"); key2 = key1
                key2_b = parse_hex_key(key1); block2_b = parse_hex_block(block2)
            else:
                key1 = coerce_key_input(key1, "Key 1"); key2 = coerce_key_input(key2, "Key 2")
                block1 = coerce_hex_input(block1, "Block 1"); block2 = block1
                key2_b = parse_hex_key(key2); block2_b = parse_hex_block(block1)
            key1_b = parse_hex_key(key1); block1_b = parse_hex_block(block1)
            rks1 = key_expansion(key1_b); rks2 = key_expansion(key2_b)
            s1 = encrypt_block_with_rounds(block1_b, rks1); s2 = encrypt_block_with_rounds(block2_b, rks2)
            mm = [bit_mismatch(s1[i], s2[i]) for i in range(11)]
            result = {
                "mode": mode, "mismatches": mm,
                "key1_hex": key1_b.hex(), "block1_hex": block1_b.hex(),
                "key2_hex": key2_b.hex(), "block2_hex": block2_b.hex(),
                "key_diff_bits": byte_bit_mismatch(key1_b, key2_b),
                "block_diff_bits": byte_bit_mismatch(block1_b, block2_b),
                "average_mismatch": round(sum(mm)/len(mm), 2),
                "final_mismatch": mm[-1], "final_percent": round(mm[-1]/128*100, 2),
                "max_mismatch": max(mm), "max_round": mm.index(max(mm)),
                "min_mismatch": min(mm), "min_round": mm.index(min(mm)),
                "closest_round": min(range(len(mm)), key=lambda i: abs(mm[i]-64)),
                "closest_mismatch": mm[min(range(len(mm)), key=lambda i: abs(mm[i]-64))],
                "round_points": [{"round": r, "mismatch": m, "percent": round(m/128*100,2),
                                   "bar_width": max(2, round(m/128*100,2)), "x": r*10,
                                   "y": round(100-(m/128*100),2)} for r,m in enumerate(mm)],
            }
        except Exception as e:
            error = str(e)
    return render_template("avalanche.html", key1=key1, block1=block1, key2=key2, block2=block2,
        mode=mode, result=result, error=error,
        key_hint="Up to 32 hex chars", block_hint="Up to 32 hex chars")

@main_bp.route("/key-expansion", methods=["GET", "POST"])
def key_expansion_page():
    error = None; result = None
    key = request.form.get("key", DEFAULT_KEY).strip()
    if request.method == "POST":
        try:
            key = coerce_key_input(key, "Key"); key_b = parse_hex_key(key); rks = key_expansion(key_b)
            result = {"key_hex": key_b.hex(), "round_keys": words_to_hex_words(rks),
                      "word_details": key_expansion_word_details(key_b)}
        except Exception as e:
            error = str(e)
    return render_template("key_expansion.html", key=key, result=result, error=error,
        key_hint="Up to 32 hex chars (short keys zero-padded on the right)")

@main_bp.route("/sbox")
def sbox():
    return render_template("sbox.html",
        sbox_grid=[[SBOX[r*16+c] for c in range(16)] for r in range(16)],
        inv_sbox_grid=[[INV_SBOX[r*16+c] for c in range(16)] for r in range(16)],
        rcon_list=list(RCON[:11]))

@main_bp.route("/coach", methods=["GET", "POST"])
def coach():
    error = None; practice_result = None; saved = False
    practice_mode  = request.form.get("practice_mode",  "encrypt")
    practice_key   = request.form.get("practice_key",   NIST_KEY).strip()
    practice_block = request.form.get("practice_block", NIST_PT).strip()
    final_answer   = request.form.get("final_answer",   "").strip()
    round_answers  = [request.form.get(f"round_answer_{r}", "").strip() for r in range(11)]
    if request.method == "POST":
        try:
            practice_result = build_coach_practice(practice_mode, practice_key,
                                                   practice_block, final_answer, round_answers)
            if current_user.is_authenticated and practice_result:
                db.session.add(CoachAttempt(
                    user_id   = current_user.id,
                    mode      = practice_result["mode"],
                    key_hex   = practice_result["key_hex"],
                    block_hex = practice_result["block_hex"],
                    correct   = practice_result["correct_count"],
                    total     = practice_result["answered_count"],
                ))
                db.session.commit(); saved = True
        except Exception as e:
            error = str(e)
    return render_template("coach.html", error=error, practice_result=practice_result,
        practice_mode=practice_mode, practice_key=practice_key, practice_block=practice_block,
        final_answer=final_answer, round_answers=round_answers, saved=saved)

@main_bp.route("/history")
@login_required
def history():
    coach_attempts = CoachAttempt.query.filter_by(
        user_id=current_user.id).order_by(CoachAttempt.created_at.desc()).limit(50).all()
    enc_attempts = EncryptionAttempt.query.filter_by(
        user_id=current_user.id).order_by(EncryptionAttempt.created_at.desc()).limit(50).all()
    return render_template("history.html", coach_attempts=coach_attempts, enc_attempts=enc_attempts)

@main_bp.route("/save-coach-attempt", methods=["POST"])
@login_required
def save_coach_attempt():
    from flask import redirect, url_for
    try:
        db.session.add(CoachAttempt(
            user_id   = current_user.id,
            mode      = request.form.get("mode", "encrypt"),
            key_hex   = request.form.get("key_hex", ""),
            block_hex = request.form.get("block_hex", ""),
            correct   = int(request.form.get("correct", 0)),
            total     = int(request.form.get("total", 0)),
        ))
        db.session.commit()
        flash("Attempt saved to your history.", "success")
    except Exception as e:
        flash(f"Could not save: {e}", "danger")
    return redirect(request.referrer or url_for("main.coach"))


@main_bp.route("/save-encryption-attempt", methods=["POST"])
@login_required
def save_encryption_attempt():
    from flask import redirect, url_for
    try:
        db.session.add(EncryptionAttempt(
            user_id    = current_user.id,
            tool_name  = request.form.get("tool_name", "unknown"),
            mode       = request.form.get("mode", "encrypt"),
            key_hex    = request.form.get("key_hex", ""),
            input_hex  = request.form.get("input_hex", ""),
            output_hex = request.form.get("output_hex", ""),
        ))
        db.session.commit()
        flash("Result saved to your history.", "success")
    except Exception as e:
        flash(f"Could not save: {e}", "danger")
    return redirect(request.referrer or url_for("main.index"))


@main_bp.route("/api/random-key")
def api_random_key():
    return jsonify({"key": secrets.token_hex(16)})

@main_bp.route("/api/random-block")
def api_random_block():
    return jsonify({"block": secrets.token_hex(16)})

@main_bp.route("/api/self-test")
def api_self_test():
    try:
        key = parse_hex_key(NIST_KEY); block = parse_hex_block(NIST_PT)
        rks = key_expansion(key); ct = encrypt_block(block, rks)
        ok = ct.hex() == NIST_CT and decrypt_block(ct, rks) == block
        return jsonify({"passed": ok, "message": "NIST vector OK" if ok else "NIST vector FAILED"})
    except Exception as e:
        return jsonify({"passed": False, "message": str(e)})

app = create_app()

if __name__ == "__main__":
    app.run(debug=True, port=5001)
